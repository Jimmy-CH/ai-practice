"""培训课件大纲转 PPTX（基于圆通品牌模板）。

职责边界：只做「Markdown 大纲 -> 可编辑 PPTX 字节」的纯转换，不读写数据库、不调用大模型。

模板：``app/assets/training_template.pptx``（10 x 5.625 英寸，16:9，品牌紫白）。
生成前清除模板自带的示例页，只复用其「封面 / 目录 / 图片页 / 封底」四个版式，
因此背景图、logo 与装饰均来自模板，产出仍可在 PowerPoint / WPS 中继续编辑。

分页规则对模型输出做兜底：``##`` 标题作为幻灯片标题，列表项作为要点，
无标题时把段落归入概述页；单页要点过多时自动拆出「（续）」页。
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Inches, Pt

TEMPLATE_PATH = Path(__file__).resolve().parent.parent / "assets" / "training_template.pptx"

FONT_NAME = "微软雅黑"
BRAND_PURPLE = RGBColor(0x4E, 0x14, 0x8B)
BRAND_ORANGE = RGBColor(0xFF, 0x69, 0x00)
TEXT_MAIN = RGBColor(0x53, 0x56, 0x5A)
TEXT_SUB = RGBColor(0x60, 0x62, 0x66)
TEXT_MUTED = RGBColor(0x90, 0x93, 0x99)

LAYOUT_COVER = "封面"
LAYOUT_AGENDA = "目录"
LAYOUT_CONTENT = "图片页"
LAYOUT_CLOSING = "封底"

MAX_BULLETS_PER_SLIDE = 6
MAX_SLIDES = 40
MAX_AGENDA_ITEMS = 8
AGENDA_TITLE_MAX = 12
RISK_KEYWORDS = ("注意", "风险")


@dataclass
class OutlineSlide:
    """一页幻灯片的中间表示；level 0 为主要点，1 为子要点。"""

    title: str
    bullets: list[tuple[int, str]] = field(default_factory=list)


_HEADING_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")
_BULLET_RE = re.compile(r"^(\s*)(?:[-*+•]|\d+[.、)])\s+(.+)$")
_TABLE_SEPARATOR_RE = re.compile(r"^\|?[\s:|-]+\|[\s:|-]*\|?$")
_LINK_RE = re.compile(r"\[([^\]]+)\]\([^)]+\)")
_STRONG_RE = re.compile(r"(\*\*|__)(.+?)\1")
_EM_RE = re.compile(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)")
_CODE_RE = re.compile(r"`([^`]*)`")
_TAG_RE = re.compile(r"<[^>]+>")
_UNSAFE_FILENAME_RE = re.compile(r'[\\/:*?"<>|\r\n]+')


def _clean_inline(text: str) -> str:
    """去掉行内 Markdown 标记，只保留可读文字。"""
    text = _LINK_RE.sub(r"\1", text)
    text = _CODE_RE.sub(r"\1", text)
    text = _STRONG_RE.sub(r"\2", text)
    text = _EM_RE.sub(r"\1", text)
    text = text.replace("~~", "").replace("**", "").replace("__", "")
    text = _TAG_RE.sub("", text)
    return text.strip()


def parse_outline(content: str, fallback_title: str = "培训课件") -> tuple[str, list[OutlineSlide]]:
    """把 Markdown 大纲解析为「文档标题 + 章节页列表」。

    兼容没有 ``##`` 的自由文本：段落与列表统一归入概述页，避免生成失败。
    """
    lines = content.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    doc_title = ""
    slides: list[OutlineSlide] = []
    current: OutlineSlide | None = None

    def start_slide(title: str) -> None:
        nonlocal current
        current = OutlineSlide(title=_clean_inline(title) or fallback_title)
        slides.append(current)

    for raw in lines:
        stripped = raw.strip()
        if not stripped or stripped.startswith("```") or stripped in {"---", "***", "___"}:
            continue

        heading = _HEADING_RE.match(stripped)
        if heading:
            level = len(heading.group(1))
            text = _clean_inline(heading.group(2))
            if not text:
                continue
            if level == 1 and not doc_title:
                doc_title = text
                continue
            if level <= 2 or current is None:
                start_slide(text)
            else:
                current.bullets.append((1, text))
            continue

        if _TABLE_SEPARATOR_RE.match(stripped):
            continue
        if stripped.startswith("|") and stripped.endswith("|"):
            cells = [_clean_inline(cell) for cell in stripped.strip("|").split("|")]
            text = " · ".join(cell for cell in cells if cell)
            if text:
                if current is None:
                    start_slide("培训要点")
                current.bullets.append((0, text))
            continue

        bullet = _BULLET_RE.match(raw)
        if bullet:
            text = _clean_inline(bullet.group(2))
            if not text:
                continue
            indent = len(bullet.group(1).replace("\t", "  "))
            if current is None:
                start_slide("培训要点")
            assert current is not None
            current.bullets.append((1 if indent >= 2 else 0, text))
            continue

        text = _clean_inline(stripped.lstrip(">").strip())
        if not text:
            continue
        if current is None:
            start_slide("培训概述")
        assert current is not None
        current.bullets.append((0, text))

    if not slides:
        start_slide("培训概述")
    return doc_title or fallback_title, slides


def split_slides(slides: list[OutlineSlide]) -> list[OutlineSlide]:
    """按「主要点 + 其子要点」成组切分，避免父子要点被拆到两页。"""
    result: list[OutlineSlide] = []
    for slide in slides:
        if not slide.bullets:
            result.append(OutlineSlide(title=slide.title))
            continue

        units: list[list[tuple[int, str]]] = []
        for level, text in slide.bullets:
            if level == 0 or not units:
                units.append([(level, text)])
            else:
                units[-1].append((level, text))

        part = 0
        chunk: list[tuple[int, str]] = []
        for unit in units:
            if chunk and len(chunk) + len(unit) > MAX_BULLETS_PER_SLIDE:
                title = slide.title if part == 0 else f"{slide.title}（续）"
                result.append(OutlineSlide(title=title, bullets=chunk))
                part += 1
                chunk = []
            chunk.extend(unit)
        if chunk:
            title = slide.title if part == 0 else f"{slide.title}（续）"
            result.append(OutlineSlide(title=title, bullets=chunk))

    if len(result) > MAX_SLIDES:
        result = result[:MAX_SLIDES]
        result.append(
            OutlineSlide(
                title="内容较多",
                bullets=[(0, "原始大纲篇幅较长，PPT 已截取前 40 页，完整内容请查看页面上的培训课件大纲。")],
            )
        )
    return result


def build_filename(title: str, audience: str = "") -> str:
    """生成安全的中文下载文件名。"""
    base = _UNSAFE_FILENAME_RE.sub("", title or "培训课件").strip()[:60] or "培训课件"
    audience_part = _UNSAFE_FILENAME_RE.sub("", audience or "").strip()
    stem = f"{base}-{audience_part}" if audience_part else base
    if not stem.endswith("培训课件"):
        stem = f"{stem}-培训课件"
    return f"{stem}.pptx"


# —— 模板与版式 ——


def _open_template() -> Presentation:
    """打开模板并清除其自带示例页，只保留母版与版式。"""
    if not TEMPLATE_PATH.exists():
        raise RuntimeError(f"PPT 模板文件缺失：{TEMPLATE_PATH}")
    prs = Presentation(str(TEMPLATE_PATH))
    slide_ids = prs.slides._sldIdLst
    for slide_id in list(slide_ids):
        prs.part.drop_rel(slide_id.rId)
        slide_ids.remove(slide_id)
    return prs


def _layout(prs: Presentation, name: str):
    for layout in prs.slide_layouts:
        if layout.name == name:
            return layout
    raise RuntimeError(f"PPT 模板缺少版式：{name}")


def _find_placeholder(slide, idx: int):
    for shape in slide.placeholders:
        if shape.placeholder_format.idx == idx:
            return shape
    return None


def _remove_placeholder(slide, idx: int) -> None:
    shape = _find_placeholder(slide, idx)
    if shape is not None:
        shape._element.getparent().remove(shape._element)


def _set_run_font(run, size: int, color: RGBColor, bold: bool = False) -> None:
    font = run.font
    font.size = Pt(size)
    font.bold = bold
    font.color.rgb = color
    font.name = FONT_NAME
    # font.name 只写 latin 字体，东亚字符需额外写 a:ea，否则 PowerPoint 会回落到主题宋体
    rpr = run._r.get_or_add_rPr()
    east_asian = rpr.find(qn("a:ea"))
    if east_asian is None:
        east_asian = parse_xml(f'<a:ea {nsdecls("a")} typeface="{FONT_NAME}"/>')
        rpr.append(east_asian)
    else:
        east_asian.set("typeface", FONT_NAME)


def _set_placeholder_text(
    shape,
    lines: list[tuple[str, int, RGBColor, bool]],
    align=PP_ALIGN.LEFT,
    line_spacing: float = 1.0,
    space_after: int = 0,
) -> None:
    """按行填充占位符，整段覆盖模板示例文字。"""
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    for index, (text, size, color, bold) in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.line_spacing = line_spacing
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(space_after)
        run = paragraph.add_run()
        run.text = text
        _set_run_font(run, size, color, bold)


def _add_text(
    slide,
    left: float,
    top: float,
    width: float,
    height: float,
    text: str,
    size: int,
    color: RGBColor = TEXT_MAIN,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    line_spacing: float = 1.0,
):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.line_spacing = line_spacing
    run = paragraph.add_run()
    run.text = text
    _set_run_font(run, size, color, bold)
    return box


def _add_rect(slide, left: float, top: float, width: float, height: float, color: RGBColor):
    rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    rect.fill.solid()
    rect.fill.fore_color.rgb = color
    rect.line.fill.background()
    rect.shadow.inherit = False
    return rect


def _fit_title_size(title: str) -> int:
    if len(title) <= 18:
        return 36
    if len(title) <= 26:
        return 32
    return 28


def _shorten(text: str, limit: int) -> str:
    return text if len(text) <= limit else f"{text[: limit - 1]}…"


# —— 版式填充 ——


def _add_cover(prs: Presentation, doc_title: str, audience: str) -> None:
    slide = prs.slides.add_slide(_layout(prs, LAYOUT_COVER))
    title_ph = _find_placeholder(slide, 11)
    if title_ph is not None:
        lines = [(doc_title, _fit_title_size(doc_title), BRAND_PURPLE, True)]
        if audience:
            lines.append((f"适用对象：{audience}", 18, TEXT_MAIN, False))
        _set_placeholder_text(title_ph, lines, line_spacing=1.15, space_after=8)

    date_ph = _find_placeholder(slide, 12)
    if date_ph is not None:
        _set_placeholder_text(
            date_ph, [(f"{date.today():%Y-%m-%d}", 16, TEXT_MUTED, False)]
        )


def _add_agenda(prs: Presentation, titles: list[str], page_no: int) -> None:
    slide = prs.slides.add_slide(_layout(prs, LAYOUT_AGENDA))
    title_ph = _find_placeholder(slide, 10)
    if title_ph is not None:
        _set_placeholder_text(title_ph, [("·课程目录", 24, BRAND_PURPLE, True)])

    # 版式自带 4 个固定间距的条目占位符，统一删除后按章节数重排，避免空占位符与 4 条上限
    for idx in (11, 12, 13, 14):
        _remove_placeholder(slide, idx)

    visible = titles[:MAX_AGENDA_ITEMS]
    size = 14 if len(visible) <= 6 else 13
    top = 1.55
    for index, title in enumerate(visible):
        y = top + index * 0.46
        _add_text(
            slide, 1.02, y, 0.42, 0.32, f"{index + 1:02d}", size, BRAND_PURPLE, bold=True
        )
        _add_text(slide, 1.46, y, 2.25, 0.32, _shorten(title, AGENDA_TITLE_MAX), size, TEXT_MAIN)
    if len(titles) > len(visible):
        _add_text(slide, 1.46, top + len(visible) * 0.46, 2.25, 0.3, "……", size, TEXT_MUTED)
    _add_text(slide, 9.05, 5.18, 0.55, 0.22, f"{page_no:02d}", 10, TEXT_MUTED, align=PP_ALIGN.RIGHT)


def _add_content_slide(
    prs: Presentation, data: OutlineSlide, page_no: int, doc_title: str
) -> None:
    slide = prs.slides.add_slide(_layout(prs, LAYOUT_CONTENT))
    risk_like = any(keyword in data.title for keyword in RISK_KEYWORDS)
    accent = BRAND_ORANGE if risk_like else BRAND_PURPLE

    # 顶部占位符位于模板页眉横线之上，用作页码 + 文档名
    kicker = _find_placeholder(slide, 10)
    if kicker is not None:
        _set_placeholder_text(
            kicker,
            [(f"{page_no:02d}  {_shorten(doc_title, 18)}", 11, TEXT_MUTED, False)],
        )

    title_size = 22 if len(data.title) <= 20 else 18
    _add_text(slide, 0.44, 0.98, 9.1, 0.5, data.title, title_size, accent, bold=True)
    _add_rect(slide, 0.45, 1.5, 0.9, 0.045, accent)

    if not data.bullets:
        _add_text(slide, 0.44, 1.75, 9.1, 0.5, "本节内容见后续页面", 16, TEXT_MUTED)
    else:
        box = slide.shapes.add_textbox(Inches(0.44), Inches(1.72), Inches(9.1), Inches(3.35))
        frame = box.text_frame
        frame.word_wrap = True
        frame.vertical_anchor = MSO_ANCHOR.TOP
        frame.margin_left = 0
        frame.margin_right = 0
        frame.margin_top = 0
        frame.margin_bottom = 0
        count = len(data.bullets)
        size = 20 if count <= 3 else 18 if count <= 5 else 16
        for index, (level, text) in enumerate(data.bullets):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.line_spacing = 1.25
            paragraph.space_before = Pt(0)
            paragraph.space_after = Pt(8 if level == 0 else 2)
            run = paragraph.add_run()
            run.text = f"{'•  ' if level == 0 else '      –  '}{text}"
            _set_run_font(
                run,
                size if level == 0 else max(size - 2, 13),
                TEXT_MAIN if level == 0 else TEXT_SUB,
            )

    _add_text(slide, 9.05, 5.18, 0.55, 0.22, f"{page_no:02d}", 10, TEXT_MUTED, align=PP_ALIGN.RIGHT)


def _add_closing(prs: Presentation, doc_title: str) -> None:
    slide = prs.slides.add_slide(_layout(prs, LAYOUT_CLOSING))
    title_ph = _find_placeholder(slide, 10)
    if title_ph is not None:
        _set_placeholder_text(title_ph, [("谢谢观看", 44, BRAND_PURPLE, True)])

    note_ph = _find_placeholder(slide, 11)
    if note_ph is not None:
        _set_placeholder_text(
            note_ph, [(f"培训课件 · {date.today():%Y-%m-%d}", 10, TEXT_MUTED, False)]
        )
    _add_text(slide, 0.65, 3.15, 4.5, 0.35, "欢迎提问与交流", 16, TEXT_MAIN)


def build_training_pptx(title: str, content: str, audience: str = "", source_title: str = "") -> bytes:
    """基于品牌模板生成培训课件 PPTX 字节；内容为空时抛 ValueError。

    ``source_title`` 保留用于接口兼容，模板封面空间有限，不再展示来源文档。
    """
    if not content or not content.strip():
        raise ValueError("培训课件大纲内容为空")

    prs = _open_template()
    doc_title, sections = parse_outline(content, title or "培训课件")
    page_sections = split_slides(sections)

    _add_cover(prs, doc_title, audience)
    page_no = 2
    if len(page_sections) >= 2:
        _add_agenda(prs, [section.title for section in sections], page_no)
        page_no += 1
    for section in page_sections:
        _add_content_slide(prs, section, page_no, doc_title)
        page_no += 1
    _add_closing(prs, doc_title)

    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()
