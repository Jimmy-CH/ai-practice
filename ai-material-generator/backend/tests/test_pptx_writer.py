import io

import pytest
from pptx import Presentation
from pptx.util import Inches

from app.services import pptx_writer

SAMPLE_OUTLINE = """# 网点管家 V3.2.0 培训课件

## 一、培训目标
- 掌握异常件一键上报
  - 熟悉上报入口与必填项
- 理解派费日结的数据口径

## 二、操作演示
1. 打开首页
2. 扫描异常件
3. 提交上报

## 三、注意事项与风险提示
> 升级至 V3.2.0 以上版本
- 派费数据存在 24 小时延迟
- 如有缺失请等待次日补数，勿重复申诉
"""


def test_parse_outline_uses_h2_as_slides():
    doc_title, slides = pptx_writer.parse_outline(SAMPLE_OUTLINE, "默认标题")
    assert doc_title == "网点管家 V3.2.0 培训课件"
    assert [slide.title for slide in slides] == ["一、培训目标", "二、操作演示", "三、注意事项与风险提示"]
    assert slides[0].bullets == [
        (0, "掌握异常件一键上报"),
        (1, "熟悉上报入口与必填项"),
        (0, "理解派费日结的数据口径"),
    ]
    assert slides[1].bullets[0] == (0, "打开首页")
    assert slides[2].bullets[0] == (0, "升级至 V3.2.0 以上版本")


def test_parse_outline_falls_back_to_paragraphs():
    doc_title, slides = pptx_writer.parse_outline("第一段内容。\n第二段内容。", "默认标题")
    assert doc_title == "默认标题"
    assert len(slides) == 1
    assert slides[0].bullets == [(0, "第一段内容。"), (0, "第二段内容。")]


def test_split_slides_keeps_children_with_parent_and_marks_continuation():
    bullets = [(0, f"要点{i}") for i in range(8)]
    slides = pptx_writer.split_slides([pptx_writer.OutlineSlide(title="章节", bullets=bullets)])
    assert [slide.title for slide in slides] == ["章节", "章节（续）"]
    assert all(slide.bullets for slide in slides)
    assert slides[0].bullets[0] == (0, "要点0")


def test_build_training_pptx_uses_brand_template():
    data = pptx_writer.build_training_pptx(
        "培训课件", SAMPLE_OUTLINE, "网点管理员", "网点管家 V3.2.0 版本迭代说明",
    )
    assert data[:2] == b"PK"

    prs = Presentation(io.BytesIO(data))
    # 模板尺寸：10 x 5.625 英寸（16:9）
    assert prs.slide_width == Inches(10.0)
    assert prs.slide_height == Inches(5.625)
    # 封面 + 目录 + 3 个章节页 + 结束页
    assert len(prs.slides) == 6
    assert prs.slides[0].slide_layout.name == "封面"
    assert prs.slides[1].slide_layout.name == "目录"
    assert [slide.slide_layout.name for slide in list(prs.slides)[2:5]] == ["图片页"] * 3
    assert prs.slides[-1].slide_layout.name == "封底"

    texts = [
        shape.text_frame.text
        for slide in prs.slides
        for shape in slide.shapes
        if shape.has_text_frame
    ]
    assert any("网点管家 V3.2.0 培训课件" in text for text in texts)
    assert any("一、培训目标" in text for text in texts)
    assert any("谢谢观看" in text for text in texts)
    # 模板示例页必须被清除，不能残留占位示例文案
    joined = "\n".join(texts)
    for leftover in ("主标题内容编辑区域", "大纲标题内容", "THANK YOU"):
        assert leftover not in joined


def test_build_training_pptx_rejects_empty_content():
    with pytest.raises(ValueError):
        pptx_writer.build_training_pptx("培训课件", "   ")


def test_build_filename_is_safe():
    assert pptx_writer.build_filename("网点:管家/V3.2.0", "网点管理员") == "网点管家V3.2.0-网点管理员-培训课件.pptx"
    assert pptx_writer.build_filename("", "") == "培训课件.pptx"
