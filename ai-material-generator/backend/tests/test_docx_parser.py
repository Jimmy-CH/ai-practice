from docx import Document
from app.services.docx_parser import parse_docx_file


def test_parse_docx_file(tmp_path):
    doc = Document()
    doc.add_paragraph("第一段")
    doc.add_paragraph("第二段")
    path = tmp_path / "test.docx"
    doc.save(path)
    assert parse_docx_file(path) == "第一段\n第二段"


def test_parse_empty_docx(tmp_path):
    doc = Document()
    path = tmp_path / "empty.docx"
    doc.save(path)
    assert parse_docx_file(path) == ""
